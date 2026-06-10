import tkinter as tk
from tkinter import filedialog, messagebox
import threading
import logging
import sys
import os
import time
import unicodedata
import psutil
import keyring

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill
from selenium import webdriver
from selenium.webdriver.chrome.options import Options
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from selenium.common.exceptions import StaleElementReferenceException
from PIL import Image, ImageDraw

# ─────────────────────────────────────────────
# Логирование
# ─────────────────────────────────────────────
logging.basicConfig(
    filename="log.txt",
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(message)s",
    encoding="utf-8",
)


# ─────────────────────────────────────────────
# Selenium‑логика
# ─────────────────────────────────────────────

def build_driver() -> webdriver.Chrome:
    options = Options()
    options.add_argument("--headless=new")
    options.add_argument("--disable-gpu")
    options.add_argument("--window-size=1920,1080")
    return webdriver.Chrome(options=options)


def login(driver, wait, email: str, password: str) -> None:
    driver.get("https://xn--52-6kcatyook.xn--80aafey1amqq.xn--d1acj3b/admin")
    time.sleep(1)
    email_input = wait.until(EC.presence_of_element_located((By.XPATH, '//input[@placeholder="E-mail"]')))
    password_input = driver.find_element(By.XPATH, '//input[@placeholder="Пароль"]')
    email_input.send_keys(email)
    password_input.send_keys(password)
    driver.find_element(By.XPATH, '//a[@data-cy="login-enter-button"]').click()
    wait.until(EC.element_to_be_clickable((By.XPATH, '//div[@data-cy="activity_order"]'))).click()
    time.sleep(1)


def normalize_text(text: str) -> str:
    return unicodedata.normalize("NFKD", text.strip().lower())


def add_child(driver, wait, event_id, event_name, event_date, event_time, child_name, child_birthdate) -> bool:
    create_btn = driver.find_element(By.XPATH, '//a[@data-qtip="Создать"]')
    driver.execute_script("arguments[0].click();", create_btn)
    time.sleep(1)

    merop_input = WebDriverWait(driver, 10).until(
        EC.presence_of_element_located((
            By.CSS_SELECTOR,
            '#win-activity_order-create input[placeholder="Выберите мероприятие"]:not([disabled])'
        ))
    )
    merop_input.clear()
    merop_input.send_keys(event_id)

    dropdown_items = WebDriverWait(driver, 10).until(
        EC.presence_of_all_elements_located((By.CSS_SELECTOR, ".x-boundlist-item"))
    )
    for option in dropdown_items:
        if option.text.strip() == event_name:
            option.click()
            break

    # Дата мероприятия
    event_date_str = event_date.strftime("%d.%m.%Y") if not pd.isna(event_date) else ""
    date_input = wait.until(EC.presence_of_element_located((By.XPATH, '//input[@placeholder="Дата"]')))
    date_input.clear()
    date_input.send_keys(event_date_str)

    # Время мероприятия
    event_time_str = (
        event_time.strftime("%H:%M") if hasattr(event_time, "strftime") else str(event_time)
    )
    time_input = wait.until(EC.presence_of_element_located((By.XPATH, '//input[contains(@title, "HH:MM")]')))
    time_input.clear()
    time_input.send_keys(event_time_str)
    time.sleep(1)

    for option in driver.find_elements(By.XPATH, '//li[contains(@class, "x-boundlist-item")]'):
        if option.is_displayed():
            option.click()
            break
    else:
        raise RuntimeError(f"Не найден видимый элемент времени для '{event_time_str}'")

    # ФИО и ДР
    birth_date_str = (
        child_birthdate.strftime("%d.%m.%Y")
        if hasattr(child_birthdate, "strftime")
        else str(child_birthdate).strip()
    )

    modal = wait.until(EC.presence_of_element_located((
        By.XPATH, '//div[contains(@class, "x-window") and contains(.,"Создать новую заявку")]'
    )))

    fio_input = modal.find_element(By.XPATH, './/input[@placeholder="Ребенок..."]')
    fio_input.clear()
    try:
        fio_input.send_keys(child_name)
    except StaleElementReferenceException:
        fio_input = wait.until(EC.presence_of_element_located((By.XPATH, './/input[@placeholder="Ребенок..."]')))
        fio_input.send_keys(child_name)
    time.sleep(1)

    quick_wait = WebDriverWait(driver, 1.5)
    try:
        li_element = quick_wait.until(EC.element_to_be_clickable((
            By.XPATH,
            f"//li[contains(@class, 'x-boundlist-item')]"
            f"//div[contains(@class, 'municipality') and contains(text(), '{birth_date_str}')]/ancestor::li"
        )))
        driver.execute_script("arguments[0].scrollIntoView(true);", li_element)
        time.sleep(0.2)
        driver.execute_script("arguments[0].click();", li_element)
        logging.info(f"[OK] {child_name} с ДР {birth_date_str} выбран из списка")
        time.sleep(1)

    except Exception:
        logging.warning(f"[WARN] {child_name} не найден в списке — открываю расширенный поиск")
        result = _advanced_search(driver, wait, modal, child_name, birth_date_str)
        if not result:
            return False

    time.sleep(2)
    
    # ИЩЕМ КНОПКУ ЗАНОВО ЧЕРЕЗ wait, А НЕ ЧЕРЕЗ старую переменную modal
    save_button = wait.until(EC.element_to_be_clickable((
        By.XPATH, 
        '//div[contains(@class, "x-window") and contains(.,"Создать новую заявку")]//a[contains(@class, "x-btn") and .//span[text()="Сохранить"]]'
    )))
    save_button.click()
    time.sleep(2)
    return True


def _advanced_search(driver, wait, modal, fio: str, birth_date_str: str) -> bool:
    try:
        time.sleep(1)
        advanced_btn = modal.find_element(By.XPATH, './/a[@data-qtip="Расширенный поиск"]')
        advanced_btn.click()
        time.sleep(1)

        parts = fio.split()
        last_name  = parts[0] if len(parts) > 0 else ""
        first_name = parts[1] if len(parts) > 1 else ""
        middle_name = parts[2] if len(parts) > 2 else ""

        wait.until(EC.presence_of_element_located((By.NAME, "last_name"))).send_keys(last_name)
        driver.find_element(By.NAME, "first_name").send_keys(first_name)
        driver.find_element(By.NAME, "patro_name").send_keys(middle_name)
        time.sleep(1)

        search_btn = driver.find_element(
            By.XPATH, '//span[text()="Поиск"]/ancestor::a[contains(@class,"x-btn")]'
        )
        driver.execute_script("arguments[0].scrollIntoView(true);", search_btn)
        time.sleep(0.3)
        driver.execute_script("arguments[0].click();", search_btn)
        time.sleep(2)

        search_window_id = "win-activity_order-safe-kid-search"
        search_window_locator = (By.ID, search_window_id)

        # Ждём — вдруг окно само закрылось
        closed = False
        for _ in range(10):
            time.sleep(0.5)
            try:
                driver.find_element(*search_window_locator)
            except Exception:
                closed = True
                break

        if closed:
            logging.info(f"[INFO] Окно поиска закрылось автоматически для {fio}")
            # save_btn = wait.until(EC.element_to_be_clickable(
            #     (By.XPATH, '//span[text()="Сохранить"]/ancestor::a[contains(@class,"x-btn")]')
            # ))
            # save_btn.click()
            # time.sleep(1)
            return True

        wait.until(EC.presence_of_element_located(
            (By.XPATH, '//div[contains(@class, "x-grid-item-container")]//table')
        ))
        rows = driver.find_elements(By.XPATH, '//div[contains(@class, "x-grid-item-container")]//table')

        for row in rows:
            try:
                fio_cell = row.find_element(By.XPATH, './/td[1]//div[contains(@class, "x-grid-cell-inner")]').text
                dr_cell  = row.find_element(By.XPATH, './/td[4]//div[contains(@class, "x-grid-cell-inner")]').text
                if normalize_text(fio_cell) == normalize_text(fio) and normalize_text(dr_cell) == normalize_text(birth_date_str):
                    check_icon = row.find_element(By.XPATH, './/i[contains(@class,"fa-check") and contains(@style,"8BC34A")]')
                    if check_icon.is_displayed():
                        row.find_element(By.XPATH, './/td[4]').click()
                        time.sleep(1)
                        return True
                    logging.warning(f"[WARN] {fio} найден, но не подтверждён — пропускаем")
            except Exception as ex:
                logging.debug(f"Ошибка при разборе строки: {ex}")

        # Ничего не нашли — закрываем оба окна
        logging.warning(f"[WARN] {fio} не найден в расширенном поиске")
        _close_by_id(driver, search_window_id)
        _close_create_window(driver)
        return False

    except Exception as ex:
        logging.error(f"Ошибка расширенного поиска: {ex}")
        return False


def _close_by_id(driver, window_id: str) -> None:
    try:
        btn = driver.find_element(
            By.XPATH, f'//div[@id="{window_id}"]//div[contains(@class, "x-tool-close")]'
        )
        driver.execute_script("arguments[0].click();", btn)
        time.sleep(0.5)
    except Exception:
        pass


def _close_create_window(driver) -> None:
    try:
        btn = driver.find_element(
            By.XPATH,
            '//div[contains(@class, "x-window") and contains(.,"Создать новую заявку на мероприятие")]'
            '//div[contains(@class, "x-tool-close")]'
        )
        driver.execute_script("arguments[0].click();", btn)
        time.sleep(0.5)
    except Exception:
        pass


def kill_chromedriver() -> None:
    for proc in psutil.process_iter():
        if "chromedriver" in proc.name().lower():
            try:
                proc.kill()
            except Exception:
                pass


# ─────────────────────────────────────────────
# GUI
# ─────────────────────────────────────────────

COLORS = {
    "bg":         "#0F1117",
    "surface":    "#1A1D27",
    "surface2":   "#22263A",
    "accent":     "#4F8EF7",
    "accent2":    "#7B61FF",
    "success":    "#3DDC84",
    "error":      "#FF5C6C",
    "warning":    "#FFB830",
    "text":       "#E8EAF6",
    "text_muted": "#6B7280",
    "border":     "#2D3147",
}

FONT_TITLE  = ("Segoe UI", 22, "bold")
FONT_LABEL  = ("Segoe UI", 10)
FONT_SMALL  = ("Segoe UI", 9)
FONT_MONO   = ("Consolas", 9)
FONT_BTN    = ("Segoe UI Semibold", 10, "bold")
FONT_COUNT  = ("Segoe UI", 32, "bold")

# Имя сервиса для хранения credentials в Windows Credential Manager
KEYRING_SERVICE = "Квантор"
KEYRING_EMAIL_KEY = "email"
KEYRING_PASS_KEY  = "password"


def _setup_clipboard_bindings(widget):
    """
    Универсальная функция: навешивает биндинги Ctrl+C/V/X/A и Shift+Insert
    на любой виджет Entry или Text через виртуальные события tkinter.
    Виртуальные события (<<Copy>>, <<Paste>>, <<Cut>>) уже обрабатываются
    tkinter стандартно — мы лишь добавляем Shift+Insert как алиас вставки
    и Ctrl+A для выделения всего.
    """
    is_text = isinstance(widget, tk.Text)

    def do_copy(e=None):
        try:
            widget.event_generate("<<Copy>>")
        except Exception:
            pass
        return "break"

    def do_paste(e=None):
        try:
            widget.event_generate("<<Paste>>")
        except Exception:
            pass
        return "break"

    def do_cut(e=None):
        try:
            widget.event_generate("<<Cut>>")
        except Exception:
            pass
        return "break"

    def do_select_all(e=None):
        if is_text:
            widget.tag_add("sel", "1.0", "end")
            widget.mark_set("insert", "end")
        else:
            widget.select_range(0, tk.END)
            widget.icursor(tk.END)
        return "break"

    # Shift+Insert — стандартная вставка в Windows
    widget.bind("<Shift-Insert>", do_paste)
    # Ctrl+A — выделить всё (tkinter не делает это по умолчанию для Entry)
    widget.bind("<Control-a>", do_select_all)
    widget.bind("<Control-A>", do_select_all)

    # На Windows иногда нужно явно пробросить Ctrl+C/V/X,
    # если стандартный механизм tkinter не срабатывает
    widget.bind("<Control-c>", do_copy)
    widget.bind("<Control-C>", do_copy)
    widget.bind("<Control-v>", do_paste)
    widget.bind("<Control-V>", do_paste)
    widget.bind("<Control-x>", do_cut)
    widget.bind("<Control-X>", do_cut)


class ModernEntry(tk.Frame):
    """Поле ввода с подписью и подсветкой границы при фокусе."""

    def __init__(self, parent, label: str, show: str = "", **kw):
        super().__init__(parent, bg=COLORS["bg"])
        tk.Label(self, text=label, bg=COLORS["bg"], fg=COLORS["text_muted"],
                 font=FONT_SMALL).pack(anchor="w", pady=(0, 3))

        self._border_frame = tk.Frame(self, bg=COLORS["border"], bd=1)
        self._border_frame.pack(fill="x")

        inner = tk.Frame(self._border_frame, bg=COLORS["surface2"])
        inner.pack(fill="x", padx=1, pady=1)

        self.var = tk.StringVar()
        self.entry = tk.Entry(
            inner,
            textvariable=self.var,
            show=show,
            bg=COLORS["surface2"],
            fg=COLORS["text"],
            insertbackground=COLORS["accent"],
            relief="flat",
            font=FONT_LABEL,
            bd=8,
            **kw
        )
        self.entry.pack(fill="x")

        # Подсветка границы при фокусе
        self.entry.bind("<FocusIn>",  lambda e: self._border_frame.config(bg=COLORS["accent"]))
        self.entry.bind("<FocusOut>", lambda e: self._border_frame.config(bg=COLORS["border"]))

        # Контекстное меню правой кнопкой
        self._context_menu = tk.Menu(self.entry, tearoff=0,
                                     bg=COLORS["surface2"], fg=COLORS["text"],
                                     activebackground=COLORS["accent"],
                                     activeforeground="white",
                                     relief="flat", bd=1)
        self._context_menu.add_command(label="Копировать",  command=self._copy)
        self._context_menu.add_command(label="Вставить",    command=self._paste)
        self._context_menu.add_command(label="Вырезать",    command=self._cut)
        self._context_menu.add_separator()
        self._context_menu.add_command(label="Выделить всё", command=self._select_all)

        self.entry.bind("<Button-3>", self._show_context_menu)

        # Горячие клавиши через универсальную функцию
        _setup_clipboard_bindings(self.entry)

    # ── методы буфера обмена ─────────────────────

    def _copy(self):
        try:
            if self.entry.selection_present():
                text = self.entry.selection_get()
                self.entry.clipboard_clear()
                self.entry.clipboard_append(text)
        except tk.TclError:
            pass

    def _paste(self):
        try:
            text = self.entry.clipboard_get()
            if self.entry.selection_present():
                self.entry.delete(tk.SEL_FIRST, tk.SEL_LAST)
            self.entry.insert(tk.INSERT, text)
        except tk.TclError:
            pass

    def _cut(self):
        try:
            if self.entry.selection_present():
                text = self.entry.selection_get()
                self.entry.clipboard_clear()
                self.entry.clipboard_append(text)
                self.entry.delete(tk.SEL_FIRST, tk.SEL_LAST)
        except tk.TclError:
            pass

    def _select_all(self):
        self.entry.select_range(0, tk.END)
        self.entry.icursor(tk.END)

    def _show_context_menu(self, event):
        try:
            self._context_menu.tk_popup(event.x_root, event.y_root)
        finally:
            self._context_menu.grab_release()

    # ── публичный интерфейс ───────────────────────

    def get(self) -> str:
        return self.var.get()

    def set(self, value: str) -> None:
        self.var.set(value)


def resource_path(relative_path):
    try:
        base_path = sys._MEIPASS
    except Exception:
        base_path = os.path.abspath(".")
    return os.path.join(base_path, relative_path)

class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("Авто-Навигатор52 - Кванториум Бор")
        self.geometry("780x700")

        # Установка иконки в заголовок окна
        icon_path = resource_path("icon.ico")
        if os.path.exists(icon_path):
            self.iconbitmap(icon_path)

        self.minsize(700, 620)
        self.configure(bg=COLORS["bg"])
        self.resizable(True, True)

        self._excel_path: str = ""
        self._df: pd.DataFrame | None = None
        self._running = False

        self._build_ui()
        self._load_credentials()

    # ── построение интерфейса ────────────────────

    def _build_ui(self):
        # Заголовок
        header = tk.Frame(self, bg=COLORS["surface"], pady=18)
        header.pack(fill="x")

        tk.Label(header, text="Авто-Навигатор52", bg=COLORS["surface"],
                 fg=COLORS["accent"], font=FONT_TITLE).pack()
        tk.Label(header, text="Автоматическое внесение участников мероприятий",
                 bg=COLORS["surface"], fg=COLORS["text_muted"],
                 font=FONT_SMALL).pack(pady=(2, 0))

        # Тело
        body = tk.Frame(self, bg=COLORS["bg"])
        body.pack(fill="both", expand=True, padx=28, pady=20)

        # Левая колонка
        left = tk.Frame(body, bg=COLORS["bg"])
        left.pack(side="left", fill="both", expand=True, padx=(0, 14))

        self._build_file_section(left)
        self._build_stats_card(left)
        self._build_credentials(left)
        self._build_run_button(left)

        # Правая колонка — лог
        right = tk.Frame(body, bg=COLORS["bg"])
        right.pack(side="right", fill="both", expand=True)
        self._build_log_panel(right)

    def _section_label(self, parent, text: str):
        f = tk.Frame(parent, bg=COLORS["bg"])
        f.pack(fill="x", pady=(14, 6))
        tk.Label(f, text=text, bg=COLORS["bg"], fg=COLORS["text_muted"],
                 font=("Segoe UI", 8, "bold")).pack(side="left")
        tk.Frame(f, bg=COLORS["border"], height=1).pack(side="left", fill="x", expand=True, padx=(8, 0))

    def _build_file_section(self, parent):
        self._section_label(parent, "ФАЙЛ ДАННЫХ")

        file_frame = tk.Frame(parent, bg=COLORS["surface2"],
                              highlightbackground=COLORS["border"],
                              highlightthickness=1, pady=10, padx=12)
        file_frame.pack(fill="x")

        top = tk.Frame(file_frame, bg=COLORS["surface2"])
        top.pack(fill="x")

        self._lbl_file = tk.Label(top, text="Файл не выбран",
                                  bg=COLORS["surface2"], fg=COLORS["text_muted"],
                                  font=FONT_SMALL, anchor="w", wraplength=280)
        self._lbl_file.pack(side="left", fill="x", expand=True)

        self._btn_browse = self._flat_button(top, "Обзор…", self._browse_file,
                                             color=COLORS["accent2"], width=9)
        self._btn_browse.pack(side="right", padx=(8, 0))

    def _build_stats_card(self, parent):
        self._section_label(parent, "СТАТИСТИКА ФАЙЛА")

        card = tk.Frame(parent, bg=COLORS["surface"],
                        highlightbackground=COLORS["border"],
                        highlightthickness=1)
        card.pack(fill="x")

        row = tk.Frame(card, bg=COLORS["surface"])
        row.pack(fill="x", padx=16, pady=14)

        for col_label, attr in [("Всего строк", "_cnt_total"),
                                 ("Успешно", "_cnt_ok"),
                                 ("Ошибок", "_cnt_err")]:
            col = tk.Frame(row, bg=COLORS["surface"])
            col.pack(side="left", expand=True)

            var = tk.StringVar(value="—")
            setattr(self, attr + "_var", var)
            color = (COLORS["text"] if col_label == "Всего строк"
                     else COLORS["success"] if col_label == "Успешно"
                     else COLORS["error"])
            tk.Label(col, textvariable=var, bg=COLORS["surface"],
                     fg=color, font=FONT_COUNT).pack()
            tk.Label(col, text=col_label, bg=COLORS["surface"],
                     fg=COLORS["text_muted"], font=FONT_SMALL).pack()

        # Прогресс-бар
        self._progress_frame = tk.Frame(card, bg=COLORS["border"], height=4)
        self._progress_frame.pack(fill="x", padx=16, pady=(0, 14))
        self._progress_frame.pack_propagate(False)

        self._progress_bar = tk.Frame(self._progress_frame, bg=COLORS["accent"], height=4)
        self._progress_bar.place(relwidth=0, relheight=1)

    def _build_credentials(self, parent):
        self._section_label(parent, "УЧЁТНЫЕ ДАННЫЕ")
        self._entry_email = ModernEntry(parent, "E-mail")
        self._entry_email.pack(fill="x", pady=(0, 8))
        self._entry_pass  = ModernEntry(parent, "Пароль", show="●")
        self._entry_pass.pack(fill="x")

        # # Подсказка + кнопка сброса сохранённых данных
        # cred_footer = tk.Frame(parent, bg=COLORS["bg"])
        # cred_footer.pack(fill="x", pady=(6, 0))

        # self._lbl_cred_hint = tk.Label(
        #     cred_footer,
        #     text="🔒 Данные сохраняются в Credential Manager",
        #     bg=COLORS["bg"], fg=COLORS["text_muted"], font=FONT_SMALL,
        # )
        # self._lbl_cred_hint.pack(side="left")

        # self._flat_button(
        #     cred_footer, "Забыть данные", self._clear_credentials,
        #     color=COLORS["surface2"]
        # ).pack(side="right")

    def _build_run_button(self, parent):
        tk.Frame(parent, bg=COLORS["bg"], height=18).pack()

        self._btn_run = tk.Button(
            parent, text="▶  ЗАПУСТИТЬ",
            bg=COLORS["accent"], fg="white",
            activebackground=COLORS["accent2"], activeforeground="white",
            font=FONT_BTN, relief="flat", bd=0,
            cursor="hand2", padx=0, pady=14,
            command=self._on_run,
        )
        self._btn_run.pack(fill="x")
        self._bind_hover(self._btn_run, COLORS["accent2"], COLORS["accent"])

    def _build_log_panel(self, parent):
        tk.Label(parent, text="ЖУРНАЛ ВЫПОЛНЕНИЯ",
                 bg=COLORS["bg"], fg=COLORS["text_muted"],
                 font=("Segoe UI", 8, "bold")).pack(anchor="w", pady=(14, 6))

        log_outer = tk.Frame(parent, bg=COLORS["border"])
        log_outer.pack(fill="both", expand=True)

        log_inner = tk.Frame(log_outer, bg=COLORS["surface"])
        log_inner.pack(fill="both", expand=True, padx=1, pady=1)

        self._log_text = tk.Text(
            log_inner, bg=COLORS["surface"], fg=COLORS["text"],
            font=FONT_MONO, relief="flat", bd=8,
            state="disabled", wrap="word",
            insertbackground=COLORS["accent"],
        )
        self._log_text.pack(fill="both", expand=True, side="left")

        sb = tk.Scrollbar(log_inner, command=self._log_text.yview,
                          bg=COLORS["surface2"], troughcolor=COLORS["bg"],
                          activebackground=COLORS["accent"])
        sb.pack(side="right", fill="y")
        self._log_text.config(yscrollcommand=sb.set)

        # Теги цветов
        self._log_text.tag_config("ok",   foreground=COLORS["success"])
        self._log_text.tag_config("err",  foreground=COLORS["error"])
        self._log_text.tag_config("warn", foreground=COLORS["warning"])
        self._log_text.tag_config("info", foreground=COLORS["text_muted"])

        # Контекстное меню для лога (только чтение — только Копировать)
        log_menu = tk.Menu(self._log_text, tearoff=0,
                           bg=COLORS["surface2"], fg=COLORS["text"],
                           activebackground=COLORS["accent"],
                           activeforeground="white",
                           relief="flat", bd=1)
        log_menu.add_command(label="Копировать",    command=self._log_copy)
        log_menu.add_command(label="Выделить всё",  command=self._log_select_all)
        log_menu.add_separator()
        log_menu.add_command(label="Очистить лог",  command=self._clear_log)

        def _show_log_menu(event):
            try:
                log_menu.tk_popup(event.x_root, event.y_root)
            finally:
                log_menu.grab_release()

        self._log_text.bind("<Button-3>", _show_log_menu)

        # Ctrl+C / Ctrl+A для лога (копирование выделенного)
        def _log_copy_key(e=None):
            self._log_copy()
            return "break"

        def _log_select_all_key(e=None):
            self._log_select_all()
            return "break"

        self._log_text.bind("<Control-c>", _log_copy_key)
        self._log_text.bind("<Control-C>", _log_copy_key)
        self._log_text.bind("<Control-a>", _log_select_all_key)
        self._log_text.bind("<Control-A>", _log_select_all_key)

        # Кнопка очистки
        self._flat_button(parent, "Очистить лог", self._clear_log,
                          color=COLORS["border"]).pack(anchor="e", pady=(6, 0))

    # ── вспомогательные методы UI ────────────────

    def _flat_button(self, parent, text, command, color=None, width=None):
        color = color or COLORS["surface2"]
        kw = dict(width=width) if width else {}
        btn = tk.Button(parent, text=text, command=command,
                        bg=color, fg=COLORS["text"],
                        activebackground=COLORS["accent"],
                        activeforeground="white",
                        relief="flat", bd=0, pady=6, padx=10,
                        font=FONT_SMALL, cursor="hand2", **kw)
        self._bind_hover(btn, COLORS["accent"], color)
        return btn

    @staticmethod
    def _bind_hover(widget, hover_color, normal_color):
        widget.bind("<Enter>", lambda e: widget.config(bg=hover_color))
        widget.bind("<Leave>", lambda e: widget.config(bg=normal_color))

    def _log(self, message: str, tag: str = ""):
        def _append():
            self._log_text.config(state="normal")
            prefix = {"ok": "✔ ", "err": "✘ ", "warn": "⚠ ", "info": "  "}.get(tag, "")
            self._log_text.insert("end", prefix + message + "\n", tag)
            self._log_text.see("end")
            self._log_text.config(state="disabled")
        self.after(0, _append)
        logging.info(message)

    def _clear_log(self):
        self._log_text.config(state="normal")
        self._log_text.delete("1.0", "end")
        self._log_text.config(state="disabled")

    def _log_copy(self):
        """Копировать выделенный текст из лога."""
        try:
            text = self._log_text.get(tk.SEL_FIRST, tk.SEL_LAST)
            self.clipboard_clear()
            self.clipboard_append(text)
        except tk.TclError:
            # Ничего не выделено — копируем всё
            text = self._log_text.get("1.0", "end-1c")
            self.clipboard_clear()
            self.clipboard_append(text)

    def _log_select_all(self):
        """Выделить весь текст в логе."""
        self._log_text.tag_add("sel", "1.0", "end")

    def _set_stat(self, attr, value):
        self.after(0, lambda: getattr(self, attr + "_var").set(str(value)))

    def _set_progress(self, done: int, total: int):
        ratio = done / total if total else 0
        self.after(0, lambda: self._progress_bar.place(relwidth=ratio, relheight=1))

    # ── сохранение / загрузка учётных данных ────

    def _load_credentials(self):
        """Загружает email и пароль из Windows Credential Manager при старте."""
        try:
            email = keyring.get_password(KEYRING_SERVICE, KEYRING_EMAIL_KEY)
            password = keyring.get_password(KEYRING_SERVICE, KEYRING_PASS_KEY)
            if email:
                self._entry_email.set(email)
            if password:
                self._entry_pass.set(password)
            if email or password:
                self._log("Учётные данные загружены из системного хранилища", "info")
        except Exception as ex:
            logging.warning(f"Не удалось загрузить credentials: {ex}")

    def _save_credentials(self, email: str, password: str):
        """Сохраняет email и пароль в Windows Credential Manager."""
        try:
            keyring.set_password(KEYRING_SERVICE, KEYRING_EMAIL_KEY, email)
            keyring.set_password(KEYRING_SERVICE, KEYRING_PASS_KEY, password)
            logging.info("Учётные данные сохранены в системном хранилище")
        except Exception as ex:
            logging.warning(f"Не удалось сохранить credentials: {ex}")

    def _clear_credentials(self):
        """Удаляет сохранённые данные из Credential Manager."""
        try:
            keyring.delete_password(KEYRING_SERVICE, KEYRING_EMAIL_KEY)
            keyring.delete_password(KEYRING_SERVICE, KEYRING_PASS_KEY)
            self._entry_email.set("")
            self._entry_pass.set("")
            self._log("Сохранённые учётные данные удалены", "warn")
        except keyring.errors.PasswordDeleteError:
            self._log("Сохранённых данных не найдено", "info")
        except Exception as ex:
            logging.warning(f"Ошибка при удалении credentials: {ex}")

    # ── действия ────────────────────────────────

    def _browse_file(self):
        path = filedialog.askopenfilename(
            title="Выберите Excel-файл",
            filetypes=[("Excel файлы", "*.xlsx *.xls"), ("Все файлы", "*.*")]
        )
        if not path:
            return
        self._excel_path = path
        try:
            self._df = pd.read_excel(path)
            count = len(self._df)
            self._lbl_file.config(text=os.path.basename(path), fg=COLORS["text"])
            self._set_stat("_cnt_total", count)
            self._set_stat("_cnt_ok", "—")
            self._set_stat("_cnt_err", "—")
            self._set_progress(0, 1)
            self._log(f"Загружен файл: {os.path.basename(path)} ({count} записей)", "info")
        except Exception as ex:
            messagebox.showerror("Ошибка", f"Не удалось прочитать файл:\n{ex}")

    def _on_run(self):
        if self._running:
            return

        if not self._excel_path or self._df is None:
            messagebox.showwarning("Нет файла", "Сначала выберите Excel-файл.")
            return

        email = self._entry_email.get().strip()
        password = self._entry_pass.get().strip()
        if not email or not password:
            messagebox.showwarning("Нет данных", "Введите E-mail и пароль.")
            return

        # Сохраняем credentials в Windows Credential Manager
        self._save_credentials(email, password)

        self._running = True
        self._btn_run.config(text="⏳  Выполняется…", state="disabled",
                             bg=COLORS["surface2"])
        self._set_stat("_cnt_ok", 0)
        self._set_stat("_cnt_err", 0)
        self._set_progress(0, len(self._df))

        thread = threading.Thread(target=self._run_worker,
                                  args=(email, password), daemon=True)
        thread.start()

    def _run_worker(self, email: str, password: str):
        df = self._df.copy()
        excel_path = self._excel_path

        try:
            wb = load_workbook(excel_path)
            ws = wb.active
            red_fill = PatternFill(start_color="FFFF9999", end_color="FFFF9999", fill_type="solid")

            driver = build_driver()
            wait   = WebDriverWait(driver, 15)

            self._log("Выполняется вход на сайт…", "info")
            login(driver, wait, email, password)
            self._log("Вход выполнен успешно", "ok")

            total = len(df)
            ok_count  = 0
            err_count = 0

            for i, row in df.iterrows():
                child = row["ФИО"]
                self._log(f"Обработка [{i+1}/{total}]: {child}", "info")

                try:
                    success = add_child(
                        driver, wait,
                        row["ID"], row["Мероприятие"],
                        row["Дата МП"], row["Время МП"],
                        child, row["Дата рождения"],
                    )
                except Exception as ex:
                    logging.error(f"Ошибка add_child для {child}: {ex}")
                    success = False

                if success:
                    ok_count += 1
                    self._log(f"{child} — внесён", "ok")
                else:
                    err_count += 1
                    self._log(f"{child} — ошибка, строка помечена красным", "err")
                    for cell in ws[i + 2]:
                        cell.fill = red_fill

                self._set_stat("_cnt_ok",  ok_count)
                self._set_stat("_cnt_err", err_count)
                self._set_progress(ok_count + err_count, total)

            wb.save(excel_path)
            driver.quit()
            kill_chromedriver()

            summary = f"Готово! Успешно: {ok_count}, ошибок: {err_count}"
            self._log(summary, "ok")
            self.after(0, lambda: messagebox.showinfo("Готово", summary))

        except Exception as ex:
            logging.exception("Критическая ошибка:")
            self._log(f"Критическая ошибка: {ex}", "err")
            self.after(0, lambda: messagebox.showerror("Ошибка", str(ex)))

        finally:
            self._running = False
            self.after(0, lambda: self._btn_run.config(
                text="▶  ЗАПУСТИТЬ", state="normal", bg=COLORS["accent"]
            ))


# ─────────────────────────────────────────────
# Точка входа
# ─────────────────────────────────────────────
if __name__ == "__main__":
    app = App()
    app.mainloop()
